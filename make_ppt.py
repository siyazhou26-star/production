"""Generate editable 16:9 lab meeting slides from UTF-8 JSON."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def export(source, destination):
    data = json.loads(Path(source).read_text(encoding='utf-8'))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    pages = [{'title': data['title'], 'bullets': [data.get('subtitle', '组会汇报')]}] + data['slides']
    for number, page in enumerate(pages, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(247, 249, 252)
        title = slide.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.9), Inches(1.1)).text_frame
        title.text = page['title']
        title.paragraphs[0].font.size = Pt(32)
        title.paragraphs[0].font.bold = True
        title.paragraphs[0].font.color.rgb = RGBColor(24, 52, 88)
        body = slide.shapes.add_textbox(Inches(.8), Inches(1.9), Inches(11.7), Inches(4.8)).text_frame
        body.word_wrap = True
        for i, bullet in enumerate(page.get('bullets', [])):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(23)
            paragraph.space_after = Pt(20)
        footer = slide.shapes.add_textbox(Inches(.8), Inches(7), Inches(11.7), Inches(.3)).text_frame
        footer.text = f'组会汇报 · {number:02d}'
        footer.paragraphs[0].font.size = Pt(11)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = 'Microsoft YaHei'
        if page.get('notes'):
            slide.notes_slide.notes_text_frame.text = page['notes']
    target = Path(destination)
    target.parent.mkdir(parents=True, exist_ok=True)
    deck.save(target)
    return target


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='生成可编辑的组会汇报 PPT')
    parser.add_argument('source', help='UTF-8 JSON 内容文件')
    parser.add_argument('output', help='输出 .pptx 路径')
    args = parser.parse_args()
    print(export(args.source, args.output))
