import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation

class ExportTest(unittest.TestCase):
    def test_exports_editable_chinese_slides(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'report.pptx'
            result = subprocess.run([sys.executable, 'make_ppt.py', 'examples/report.json', str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            deck = Presentation(output)
            self.assertEqual(len(deck.slides), 7)
            texts = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn('研究背景', texts)
            self.assertIn('待填写', texts)
            self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=2)
